import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def create_support_letter():
    doc = docx.Document()

    # Page setup - Standard Letter with 1-inch margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)

    # Base Normal Style setup
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(0x11, 0x18, 0x27) # Dark slate/near-black

    # Helper function for adding paragraphs with controlled spacing
    def add_p(text="", space_after=6, line_spacing=1.15, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.LEFT):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = line_spacing
        if text:
            run = p.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
            run.bold = bold
            run.italic = italic
            run.font.color.rgb = RGBColor(0x11, 0x18, 0x27)
        return p

    # Date header
    add_p("Date: September 24, 2026", space_after=18, bold=True)

    # Salutation
    add_p("To Whom It May Concern:", space_after=10, bold=True)

    # Subject / RE line
    p_re = add_p("RE: Letter of Support for Lilia Judith Campos Martinez (Immigration Matter)", space_after=14, bold=True)

    # Body Paragraphs
    paragraphs = [
        "My name is [Your Full Name], and I am writing this letter in support of Lilia Judith Campos Martinez regarding her immigration matter.",
        "I have known Lilia since 2012, and throughout these many years, she has become an important person in my life. Lilia is a friend of my mother, but my relationship with her has grown into something much deeper. I have had the opportunity to live under her roof, spend significant time with her and her family, and personally witness the kind of person she is. Because of this, I feel comfortable speaking about her character and the positive influence she has had on my life.",
        "Lilia is a hardworking, honest, responsible, and compassionate person. She has always been someone who is willing to help others, even when she has responsibilities of her own. I have seen her care for the people around her and offer her support and guidance when someone needs it. She treats people with respect and kindness, and she has always been someone I can trust.",
        "One of the things I especially appreciate about Lilia is the guidance she gave me when I was younger. She would give me advice about making responsible decisions, working hard, and taking pride in being a responsible and dependable worker. Her advice was meaningful to me because it came from someone who practiced those same values in her own life. She taught me through both her words and her example that working hard, being responsible, and doing your best are important qualities.",
        "Lilia is also a devoted mother to her three children. I have personally witnessed the sacrifices she makes for her family and how hard she works to provide for them. Her children are very important to her, and she has always demonstrated love, dedication, and responsibility as a mother. She has worked hard to support her family and has always tried to make sure her children are cared for.",
        "Over the years, Lilia has also become like a mother to me. She has guided me, supported me, and given me advice during important times in my life. I am grateful for the role she has played in my life and for the positive example she has given me. I do not see her simply as my mother’s friend; I see her as someone who has been part of my family and who has had a meaningful influence on the person I have become.",
        "Everything I have written in this letter is based on my own personal experiences and the many years I have known Lilia. From what I have personally witnessed, she is a hardworking, honest, caring, responsible, and family-oriented person who has always tried to help others and provide for the people she loves.",
        "I respectfully submit this letter in support of Lilia Judith Campos Martinez and hope that my personal experience and knowledge of her character will be taken into consideration in her immigration matter.",
        "I am willing to confirm the information in this letter and provide additional information about my relationship with Lilia if requested.",
        "Thank you for taking the time to read my letter and consider my personal experience with Lilia."
    ]

    for p_text in paragraphs:
        add_p(p_text, space_after=10, line_spacing=1.15)

    # Closing
    add_p("Sincerely,", space_after=24)

    # Signature Block
    add_p("Signature: _________________________________________", space_after=12)
    add_p("Full Legal Name: [Your Full Legal Name]", space_after=10)
    add_p("Date: __________________", space_after=10)
    add_p("Phone Number: __________________________", space_after=10)
    add_p("Email Address: __________________________", space_after=10)
    add_p("Address: ____________________________________________________", space_after=6)

    # Output path
    output_dir = os.path.join(os.getcwd(), 'docs')
    os.makedirs(output_dir, exist_ok=True)
    output_file = os.path.join(output_dir, 'Letter_of_Support_Lilia_Judith_Campos_Martinez.docx')

    doc.save(output_file)
    print(f"[SUCCESS] Word document generated successfully at: {output_file}")
    return output_file

if __name__ == '__main__':
    create_support_letter()
